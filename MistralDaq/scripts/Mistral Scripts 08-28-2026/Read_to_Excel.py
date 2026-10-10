from maq20 import MAQ20
from maq20.modules.diol import DIOL
import time
import datetime
import numpy as numpy
from pymodbus.client import ModbusTcpClient # using pymodbus 3.9.2
from openpyxl import Workbook
def get_bit(value, bit_index):
    return (value >> bit_index) & 1
def get_bits(value, start_bit, bit_length):
    mask = (1 << bit_length) - 1
    return (value >> start_bit) & mask

startrun = datetime.datetime.now()
now = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
filenm = "./excel/CO2Chiller_" + now + ".xlsx"
wb = Workbook()
ws = wb.active
ws['A1'] = "datetime"
ws['B1'] = "i"
ws['C1'] = "KTC Temp"
ws['D1'] = "TTC Temp"
ws['E1'] = "VFD Freq"
ws['F1'] = "SYS4540"
ws['G1'] = "Water Flow"
ws['H1'] = "Water Temp"
ws['I1'] = "Mass Flow"
ws['J1'] = "Chiller_Ok"
ws['K1'] = "HC_Ena"
ws['L1'] = "Pump_Ena"
ws['M1'] = "Heater_Ena"
ws['N1'] = "Flow_Error"
ws['O1'] = "Fire_Alarm"
ws['P1'] = "MEV_Ena"
ws['Q1'] = "ACC_PT1"
ws['R1'] = "PumpIn_PT2"
ws['S1'] = "PumpOut_PT3"
ws['T1'] = "FM_OUT_PT4"
ws['U1'] = "P499_PT5"
ws['V1'] = "Heater Imon"
ws['W1'] = "suction T"
ws['X1'] = "Dischrg T"
ws['Y1'] = "Liquid T"
ws['Z1'] = "Evap_In"
ws['AA1'] = "Evap_Out"
ws['AB1'] = "Water_In"
ws['AC1'] = "Water_Out"
ws['AD1'] = "Pump_In"
ws['AE1'] = "Nolabel1"
ws['AF1'] = "Nolabel2"
ws['AG1'] = "Hex_In"
ws['AH1'] = "Hex_Out"
ws['AI1'] = "Nolabel3"

wb.save(filenm)
   
i=0


maq20 = MAQ20(ip_address="192.168.1.40", port=502)
isn1  = maq20.find("S0157858-08")
diol1 = maq20.find("S0107173-06")
vout  = maq20.find("S0156303-06")
ttc1  = maq20.find("S0170586-07")
ttc2  = maq20.find("S0176524-04")
#print(isn1)
#print(diol1)
#print(vout)
#print(ttc1)
#print(ttc2)
#CH0_ttc = ttc1.read_channel_data(0)
#ttc_array = ttc1.read_data(start_Channel = 0, number_of_channels=4)
#diol1.write_channel_data(channel=4, data=0)
#iout.write_channel_data(channel=0,data=12.5)
#vout.write_channel_data(channel=0,data=0)

#isn1_data=isn1.read_data(0,16)
#ttc1_data=ttc1.read_data(0,8)
#ttc2_data=ttc2.read_data(0,8)
#print(isn1_data)
#print(ttc1_data)
#print(ttc2_data)
#print(ttc2_data[2])

Gateway = ModbusTcpClient('192.168.1.50')      #default is port 502
while(i<1000):
        
    Gateway.connect()

    regs_l = Gateway.read_holding_registers(address = 0 , count = 125, device_id = 36) #read PLC registers
    #print(regs_l)

    CH0_KTC_Calc = ((((regs_l.registers[67]-4000)/1000)*10)-60)      #maps -60C = 4mA, 0C = 10mA, 100C = 20mA, 10C/mA
    CH0_KTC_mapped = regs_l.registers[68]                         #is a signed number, 
    CH1_TTC_Calc = (((regs_l.registers[69]-4000)/1000)*10)-60
    CH1_TTC_mapped = regs_l.registers[70]
    CH2_VFD_F_Calc = (((regs_l.registers[71])+150)/564)            # GoogleSheets Yaskawa VFD read 
    CH2_VFD_F_mapped = (regs_l.registers[72])
    CH3_System450_Calc = (regs_l.registers[73])/10
    CH3_System450_mapped = (regs_l.registers[74])
    CH4_Water_Flow_Calc = (((regs_l.registers[75]-4000)/1000)*412.5)/1000
    CH4_Water_Flow_mapped = (regs_l.registers[76]/1000)
    CH5_Water_Temp_Calc = (((regs_l.registers[77]-4000)/1000)*6.25)-20   #maps -20C = 4mA,80C = 20mA, 6.25C/mA
    CH5_Water_Temp_mapped = (regs_l.registers[78])
    CH6_RHE27_Iout_Calc = (((regs_l.registers[79]-4000)/1000)*62.5)
    #CH6_f_MFlow = f"{CH6_RHE27_Iout_Calc:.1f}"
    CH6_RHE27_Iout_mapped = (regs_l.registers[80]/1000)

    #print('\nKTC_CH0 reg: {} CH0 Tcalc: {:.2f}, CH0 T: {}'.format(regs_l.registers[67],CH0_KTC_Current,CH0_KTC_Temp))
    """
    print('CH0_KTC_Calc {:.2f}'.format(CH0_KTC_Calc))
    print('CH0_KTC_mapped',CH0_KTC_mapped)
    print('CH1_TTC_Calc {:.2f}'.format(CH1_TTC_Calc))
    print('CH1_TTC_mapped',CH1_TTC_mapped)
    print('CH2_VFD_F_Calc {:.2f}'.format(CH2_VFD_F_Calc))
    print('CH2_VFD_F_mapped',CH2_VFD_F_mapped)
    print('CH3_System450_Calc {:.2f}'.format(CH3_System450_Calc))
    print('CH3_System450_mapped',CH3_System450_mapped)
    print('CH4_Water_Flow_Calc {:.2f}'.format(CH4_Water_Flow_Calc))
    print('CH4_Water_Flow_mapped',CH4_Water_Flow_mapped)
    print('CH5_Water_Temp_Calc {:.2f}'.format(CH5_Water_Temp_Calc))
    print('CH5_Water_Temp_mapped',CH5_Water_Temp_mapped)
    print('CH6_RHE27_Iout_Calc {:.2f}'.format(CH6_RHE27_Iout_Calc))
    print('CH6_RHE27_Iout_mapped',CH6_RHE27_Iout_mapped)
    """
    #print('TTC_CH2 reg: {} CH2 Tcalc: {:.2f}, CH2 T: {}'.format(regs_l.registers[71],CH2_TTC_Current,CH2_TTC_Temp))
    #print('Vac = {},Torr = {}' .format(PLC_VAC_Volts,Vacuum_Pressure))

    Gate10 = (regs_l.registers[86])
    Error_reg = (regs_l.registers[87])
    Gate11 = (regs_l.registers[88])     
    Gate12 = (regs_l.registers[90])
    Gate13 = (regs_l.registers[92])
    #print("Gate10 = {:b}".format(Gate10))
    #print("i=",i)
    print("i={}, Gate11 = {:b}".format(i,Gate11))
    #print("Gate12 = {:b}".format(Gate12))
    #print("Gate13 = {:b}".format(Gate13))


    """
    ToGate11=ToGateway_User_B(P(/SM_10Hz),10,Master_Reset,RHE27_DO0,HC_Status,VFD_Status,Fire_Alarm,Chiller_Ok,HC_Ok,Flow_Error_LatchOK,0)
    ToGate12=ToGateway_User_B(P(/ToGate10),11,KTC_CH0_Valid,TTC_CH1_Valid,VFD_ReadF_Valid,Sys450_Read_Valid,Water_Flow_Valid,Water_Temp_Valid,RHE27_Valid,0,SR_ErrorCode)
    ToGate13=ToGateway_User_C(P(/ToGate11),12,Heater_Ena,RHE27_DI1,Heater_Ena_LED,Fire_Alarm_LED,Drive_Run_Relay,Drive_Reset_Relay,Drive_Sporlan_Relay,Chiller_Ok_LED,HC_Ena_LED,Flow_Error_LED,Pump_Ena,KTC_Temp_Ok,TTC_Temp_Ok,HeatCraft_Enable,Run_Enable,Stop_Pump,0)
    ToGate13=ToGateway_User_C(P(/ToGate11),12,KTC_Temp_Filter,TTC_Temp_Filter,RHE27_Flow_Filter,HC_Status_Filter,Water_Flow_Ok_Filter,Water_Temp_Ok_Filter,0,0,0,0,0,0,0,0,0,0,0)
    """

    Chiller_Ok_LED = get_bit(Gate11,0)                    
    HC_Ena_LED = get_bit(Gate11,1)                    
    Pump_Ena_LED = get_bit(Gate11,2)                    
    Heater_Ena_LED = get_bit(Gate11,3)
    Flow_Error_LED = get_bit(Gate11,4)
    Fire_Alarm_LED = get_bit(Gate11,5)
    MEV_Ena = get_bit(Gate11,6)
    
    #maq20 = MAQ20(ip_address="192.168.1.40", port=502)
    #time.sleep(.2)
    try:
        isn1_data=isn1.read_data(0,16)
        
    except:
        maq20 = MAQ20(ip_address="192.168.1.40", port=502)
        isn1  = maq20.find("S0157858-08")
        diol1 = maq20.find("S0107173-06")
        vout  = maq20.find("S0156303-06")
        ttc1  = maq20.find("S0170586-07")
        ttc2  = maq20.find("S0176524-04")
        time.sleep(2)
        print("Read isn1 error")
        pass

    try:
        ttc1_data=ttc1.read_data(0,8)
    except:
        maq20 = MAQ20(ip_address="192.168.1.40", port=502)
        isn1  = maq20.find("S0157858-08")
        diol1 = maq20.find("S0107173-06")
        vout  = maq20.find("S0156303-06")
        ttc1  = maq20.find("S0170586-07")
        ttc2  = maq20.find("S0176524-04")        
        time.sleep(2)
        print("ttc1 Read error")
        pass
    try:
        ttc2_data=ttc1.read_data(0,8)
    except:
        maq20 = MAQ20(ip_address="192.168.1.40", port=502)
        isn1  = maq20.find("S0157858-08")
        diol1 = maq20.find("S0107173-06")
        vout  = maq20.find("S0156303-06")
        ttc1  = maq20.find("S0170586-07")
        ttc2  = maq20.find("S0176524-04")        
        time.sleep(2)
        print("ttc2 Read error")
        pass
                
    
   

    ws.append([datetime.datetime.now(),i,CH0_KTC_Calc,CH1_TTC_Calc,CH2_VFD_F_Calc,CH3_System450_Calc,CH4_Water_Flow_Calc,CH5_Water_Temp_Calc,CH6_RHE27_Iout_Calc,Chiller_Ok_LED,HC_Ena_LED,Pump_Ena_LED,Heater_Ena_LED,Flow_Error_LED,Fire_Alarm_LED,MEV_Ena,isn1_data[0], isn1_data[1],isn1_data[2],isn1_data[7],isn1_data[8],isn1_data[7],ttc1_data[0],ttc1_data[1],ttc1_data[2],ttc1_data[3],ttc1_data[4],ttc1_data[5],ttc1_data[6],ttc2_data[0],ttc2_data[1],ttc2_data[2],ttc2_data[3],ttc2_data[4],ttc2_data[5]])
    wb.save(filenm)
    i+=1
    Gateway.close()
    time.sleep(10)
